"""Read-only workbook extraction. Writes a private, reviewable local import preview."""
import hashlib, json, sys, uuid
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, data_only=True)
def cents(value):
    return int((Decimal(str(value or 0))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def string(value): return str(value or '').strip()
def date(value): return value.strftime('%Y-%m-%d') if isinstance(value,datetime) else ''
def uid(kind,key): return str(uuid.uuid5(uuid.NAMESPACE_URL,f'prima-import:{kind}:{key}'))
fingerprint=hashlib.sha256(source.read_bytes()).hexdigest()
warnings=['Historical received-to-date amounts are imported as one receipt per order. Payment dates are unknown; no installment history is invented.', 'Production/design statuses are not present in the workbook. Imported orders are marked Needs review / Pending artwork.', 'Partner balances include only paid partner-funded expenses. Some workbook capital formulas may also include unpaid amounts.', 'Monthly receipts use actual payment dates. Undated imported receipts appear only in all-time totals.']
blockers=[]
items={}
for row in book['Order Items'].iter_rows(min_row=5,max_row=1004,values_only=True):
    if not isinstance(row[0],(int,float)): continue
    key=int(row[0]); name=string(row[1]); quantity=row[2]; rate=row[3]
    if not name or not isinstance(quantity,(int,float)) or quantity<=0 or int(quantity)!=quantity or not isinstance(rate,(int,float)) or rate<0:
        blockers.append(f'Invalid item on legacy order {key}'); continue
    items.setdefault(key,[]).append(dict(id=uid('item',f'{key}:{len(items.get(key,[]))}'),name=name,category=name,specification='',quantity=int(quantity),unitPrice=cents(rate)))
orders=[];payments=[];expenses=[];ids=set()
for index,row in enumerate(book['Orders Log'].iter_rows(min_row=5,max_row=199,values_only=True),5):
    if not isinstance(row[0],(int,float)): continue
    key=int(row[0]); oid=uid('order',key)
    if key in ids: blockers.append(f'Duplicate order ID {key}'); continue
    ids.add(key)
    if not items.get(key): blockers.append(f'Order {key} has no product lines')
    orders.append(dict(id=oid,number=f'XL-{key:05}',date=date(row[1]),customer=string(row[2]),brand=string(row[3]),phone='',address='',source='Other',status='Needs review',design='Pending artwork',dueDate='',notes=f'Imported from Orders Log row {index}. '+string(row[11]),artwork='',items=items.get(key,[]),discount=0,deliveryCharge=0,version=1,createdAt=datetime.now(timezone.utc).isoformat()))
    if cents(row[8]): payments.append(dict(id=uid('receipt',key),orderId=oid,amount=cents(row[8]),date='',method='Historical aggregate',account='Business',reference=f'Excel order {key}',kind='Receipt',note='Received to date including advance. Actual payment dates were not recorded.'))
    if cents(row[12]):
        funding=string(row[13]) or 'Business'; paid=string(row[14]).lower()=='yes'
        if funding not in ['Business','Ismail','Rizwan']: blockers.append(f'Unknown delivery funding source: {funding}')
        expenses.append(dict(id=uid('delivery',key),date=date(row[1]),description=f'Delivery — {string(row[3]) or string(row[2])}',category='Delivery',productType='Other',brand=string(row[3]),orderId=oid,amount=cents(row[12]),funding=funding,paid=paid,paidDate='',note=f'Imported once from Orders Log row {index}; exact payment date unknown.'))
unlinked=0
for index,row in enumerate(book['Expenses Log'].iter_rows(min_row=5,max_row=198,values_only=True),5):
    if not isinstance(row[0],(int,float)): continue
    brand=string(row[3]); matches=[o for o in orders if brand and brand!='-' and o['brand'].casefold()==brand.casefold()]
    oid=matches[0]['id'] if len(matches)==1 else ''
    if not oid and string(row[4]).lower()=='product': unlinked+=1
    funding=string(row[5]); paid=string(row[8]).lower()=='yes'
    if funding not in ['Business','Ismail','Rizwan']: blockers.append(f'Unknown expense funding source at row {index}: {funding}')
    category=string(row[4]).title()
    if category not in ['Product','Delivery','Marketing','Business','Other']: category='Other'
    expenses.append(dict(id=uid('expense',index),date=date(row[1]),description=string(row[2]),category=category,productType=string(row[9]) or 'Other',brand=brand,orderId=oid,amount=cents(row[6]),funding=funding,paid=paid,paidDate='',note=f'Expenses Log row {index}. '+string(row[7])))
for key in items:
    if key not in ids: blockers.append(f'Items reference missing order {key}')
if unlinked: warnings.append(f'{unlinked} product expenses could not be uniquely linked to an order by brand. They remain in overall/product costs; order cost detail needs review.')
sales=sum(i['quantity']*i['unitPrice'] for o in orders for i in o['items']); received=sum(p['amount'] for p in payments)
preview=dict(fingerprint=fingerprint,source=source.name,extractedAt=datetime.now(timezone.utc).isoformat(),orders=orders,payments=payments,expenses=expenses,openingBalance=cents(book['Business Account']['B4'].value),warnings=warnings+blockers,blockers=blockers,totals=dict(sales=sales,received=received,outstanding=sales-received,expenses=sum(e['amount'] for e in expenses),delivery=sum(e['amount'] for e in expenses if e['category']=='Delivery')))
target=Path('.local/admin/import-preview.json');target.parent.mkdir(parents=True,exist_ok=True);target.write_text(json.dumps(preview,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(dict(orders=len(orders),items=sum(len(o['items']) for o in orders),payments=len(payments),expenses=len(expenses),warnings=len(preview['warnings']),blockers=blockers)))
