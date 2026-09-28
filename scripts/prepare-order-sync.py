"""Read an updated workbook by column names; never edit the source workbook."""
import sys,json,hashlib,uuid
from pathlib import Path
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

source=Path(sys.argv[1])
book=openpyxl.load_workbook(source,read_only=True,data_only=True)
def cents(v):
    n=Decimal(str(v or 0))
    if not n.is_finite() or n<0: raise ValueError('Invalid negative/nonfinite amount')
    return int((n*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def uid(kind,key): return str(uuid.uuid5(uuid.NAMESPACE_URL,f'prima-import:{kind}:{key}'))
def rows(sheet):
    ws=book[sheet]; headers=[str(c.value or '').strip() for c in ws[4]]
    for index,row in enumerate(ws.iter_rows(min_row=5,values_only=True),5):
        if row and isinstance(row[0],(int,float)):
            yield index,dict(zip(headers,row))
items={}
for index,r in rows('Order Items'):
    key=int(r['Order ID']); q=r['Quantity']; name=str(r['Item'] or '').strip()
    if not name or not isinstance(q,(int,float)) or q<=0 or int(q)!=q: raise ValueError(f'Invalid product row {index}')
    rate=cents(r['Unit Price (Rs)'])
    assert int(q)*rate==cents(r['Item Total (Rs)']),f'Item total mismatch row {index}'
    lines=items.setdefault(key,[])
    lines.append(dict(id=uid('item',f'{key}:{len(lines)}'),name=name,category=name,specification='',quantity=int(q),unitPrice=rate))
orders=[]; skipped=[];seen=set()
for index,r in rows('Orders Log'):
    key=int(r['Order ID'])
    if not r['Customer Name'] and not r['Brand Name'] and not items.get(key):
        skipped.append(key);continue
    assert key not in seen,f'Duplicate ID {key}'
    seen.add(key)
    assert isinstance(r['Date'],datetime) and r['Customer Name'] and items.get(key),f'Incomplete order {key}'
    net=sum(i['quantity']*i['unitPrice'] for i in items[key])+cents(r['Delivery Charged (Rs)'])-cents(r['Discount (Rs)'])
    assert net==cents(r['Net Order Total (Rs)']),f'Net total mismatch {key}'
    received=cents(r['Received to Date incl. Advance (Rs)'])
    assert net-received==round(float(r['Balance Due (Rs)'] or 0)*100),f'Balance mismatch {key}'
    note=str(r['Notes'] or '').strip()
    orders.append(dict(id=uid('order',key),number=f'XL-{key:05}',legacyId=key,date=r['Date'].strftime('%Y-%m-%d'),customer=str(r['Customer Name']).strip(),brand=str(r['Brand Name'] or '').strip(),items=items[key],discount=cents(r['Discount (Rs)']),deliveryCharge=cents(r['Delivery Charged (Rs)']),received=received,deliveryCost=cents(r['Delivery Cost (Rs)']),notes=note,status='Dispatched' if note.lower()=='dispatched' else None,net=net))
assert set(items)<=seen,'Orphan product rows'
result=dict(source=source.name,fingerprint=hashlib.sha256(source.read_bytes()).hexdigest(),orders=orders,skipped=skipped,totals=dict(sales=sum(o['net'] for o in orders),received=sum(o['received'] for o in orders),delivery=sum(o['deliveryCost'] for o in orders)))
out=Path('.local/admin/order-sync.json');out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(orders=len(orders),items=sum(len(o['items']) for o in orders),skipped=skipped,totals=result['totals'])))
